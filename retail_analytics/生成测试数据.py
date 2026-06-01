import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

# 创建测试数据
data = {
    '订单ID': ['ORD001', 'ORD002', 'ORD003', 'ORD004', 'ORD005', 
                'ORD006', 'ORD007', 'ORD008', 'ORD009', 'ORD010',
                'ORD011', 'ORD012', 'ORD013', 'ORD014', 'ORD015'],
    '用户ID': ['U001', 'U002', 'U003', 'U001', 'U004', 
                'U005', 'U002', 'U006', 'U007', 'U003',
                'U008', 'U004', 'U009', 'U001', 'U010'],
    '商品ID': ['P001', 'P002', 'P003', 'P004', 'P005',
                'P006', 'P007', 'P001', 'P008', 'P009',
                'P010', 'P005', 'P011', 'P012', 'P013'],
    '商品名称': ['无线蓝牙耳机', '运动休闲双肩包', '有机牛奶1L装', '纯棉短袖T恤', '智能手环运动版',
                '保温杯不锈钢', '润肤保湿面霜', '无线蓝牙耳机', '儿童益智积木', '全麦面包整箱',
                '办公转椅人体工学', '智能手环运动版', '瑜伽垫加厚防滑', '电动牙刷替换头', '笔记本电脑支架'],
    '商品分类': ['数码电子', '箱包配饰', '食品饮料', '服装鞋帽', '数码电子',
                '家居生活', '美妆护肤', '数码电子', '母婴玩具', '食品饮料',
                '家居生活', '数码电子', '运动户外', '个护清洁', '数码电子'],
    '单价': [299.00, 189.00, 15.80, 89.00, 199.00,
             69.00, 159.00, 299.00, 129.00, 45.00,
             599.00, 199.00, 89.00, 39.00, 79.00],
    '购买数量': [1, 2, 5, 3, 1, 2, 1, 1, 2, 3, 1, 1, 1, 4, 1],
    '订单金额': [299.00, 378.00, 79.00, 267.00, 199.00,
                 138.00, 159.00, 299.00, 258.00, 135.00,
                 599.00, 199.00, 89.00, 156.00, 79.00],
    '下单时间': ['2025-01-15 10:30:00', '2025-01-15 11:15:00', '2025-01-15 12:00:00', '2025-01-15 14:20:00', '2025-01-15 15:45:00',
                 '2025-01-15 16:30:00', '2025-01-16 09:10:00', '2025-01-16 10:25:00', '2025-01-16 11:40:00', '2025-01-16 13:15:00',
                 '2025-01-16 14:50:00', '2025-01-16 16:05:00', '2025-01-17 08:45:00', '2025-01-17 10:10:00', '2025-01-17 11:35:00'],
    '支付方式': ['微信支付', '支付宝', '微信支付', '信用卡', '微信支付',
                 '支付宝', '微信支付', '银联支付', '微信支付', '支付宝',
                 '信用卡', '微信支付', '支付宝', '微信支付', '银联支付'],
    '订单状态': ['已完成', '已完成', '已完成', '已发货', '已付款',
                 '已完成', '已完成', '已取消', '已完成', '已发货',
                 '已付款', '已完成', '已完成', '已发货', '已完成'],
    '用户年龄': [28, 35, 42, 28, 31, 45, 35, 29, 33, 42, 38, 31, 27, 28, 40],
    '用户性别': ['男', '女', '男', '男', '女', '男', '女', '男', '女', '男', '男', '女', '男', '男', '女'],
    '用户城市': ['北京', '上海', '广州', '北京', '深圳',
                 '成都', '上海', '杭州', '武汉', '广州',
                 '南京', '深圳', '重庆', '北京', '西安'],
    '是否会员': ['是', '否', '是', '是', '否', '是', '否', '否', '是', '是', '否', '否', '是', '是', '否'],
    '折扣率': [0.05, 0.00, 0.10, 0.08, 0.00, 0.15, 0.00, 0.00, 0.12, 0.05,
                0.00, 0.00, 0.00, 0.20, 0.00],
    '实付金额': [284.05, 378.00, 71.10, 245.64, 199.00,
                 117.30, 159.00, 299.00, 227.04, 128.25,
                 599.00, 199.00, 89.00, 124.80, 79.00]
}

df = pd.DataFrame(data)

# 创建Excel文件
wb = Workbook()
ws = wb.active
ws.title = '零售交易数据'

# 写入数据
for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

# 设置标题行格式
header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4C1D95', end_color='4C1D95', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center')

for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment

# 调整列宽
ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 20
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 10
ws.column_dimensions['G'].width = 12
ws.column_dimensions['H'].width = 12
ws.column_dimensions['I'].width = 20
ws.column_dimensions['J'].width = 12
ws.column_dimensions['K'].width = 12
ws.column_dimensions['L'].width = 10
ws.column_dimensions['M'].width = 10
ws.column_dimensions['N'].width = 12
ws.column_dimensions['O'].width = 10
ws.column_dimensions['P'].width = 10
ws.column_dimensions['Q'].width = 12

# 保存文件
wb.save('测试数据_零售交易记录.xlsx')
print('[成功] Excel文件创建成功：测试数据_零售交易记录.xlsx')
print(f'[信息] 共生成 {len(df)} 条测试数据')
print('[字段] 包含字段：订单ID, 用户ID, 商品ID, 商品名称, 商品分类, 单价, 购买数量, 订单金额, 下单时间, 支付方式, 订单状态, 用户年龄, 用户性别, 用户城市, 是否会员, 折扣率, 实付金额')